"""Text from uploaded documents: PDF, DOCX, XLSX, CSV, plain text; images pass through.

Pure parsing, no network. Spreadsheets become CSV-like text, one row per line, so the
setup module can parse team lists deterministically before any LLM call.
"""

import csv
import io
import zipfile
from dataclasses import dataclass
from typing import Literal
from xml.etree import ElementTree

from pypdf import PdfReader

DocumentKind = Literal["pdf", "docx", "xlsx", "csv", "text", "image"]
_IMAGE_TYPES = {"image/png", "image/jpeg", "image/webp", "image/gif"}
_WORD_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class UnsupportedDocument(ValueError):
    pass


@dataclass(frozen=True)
class ExtractedDocument:
    kind: DocumentKind
    text: str  # empty for images (read by the model directly)
    rows: list[list[str]] | None = None  # spreadsheets and CSV
    mime: str = ""


def _kind(filename: str, mime: str) -> DocumentKind:
    name = filename.lower()
    if mime == "application/pdf" or name.endswith(".pdf"):
        return "pdf"
    if name.endswith(".docx"):
        return "docx"
    if name.endswith((".xlsx", ".xlsm")):
        return "xlsx"
    if mime == "text/csv" or name.endswith(".csv"):
        return "csv"
    if mime in _IMAGE_TYPES or name.endswith((".png", ".jpg", ".jpeg", ".webp", ".gif")):
        return "image"
    if mime.startswith("text/") or name.endswith((".txt", ".md")):
        return "text"
    raise UnsupportedDocument(f"cannot read '{filename}' ({mime or 'unknown type'})")


def _rows_text(rows: list[list[str]]) -> str:
    return "\n".join(", ".join(cell for cell in row) for row in rows if any(row))


def _docx_text(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    paragraphs = []
    for paragraph in root.iter(f"{_WORD_NS}p"):
        paragraphs.append("".join(node.text or "" for node in paragraph.iter(f"{_WORD_NS}t")))
    return "\n".join(p for p in paragraphs if p.strip())


def _xlsx_rows(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    rows: list[list[str]] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            rows.append(["" if cell is None else str(cell).strip() for cell in row])
    return rows


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract(data: bytes, filename: str, mime: str = "") -> ExtractedDocument:
    kind = _kind(filename, mime)
    if kind == "pdf":
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
        return ExtractedDocument(kind, text.strip(), mime=mime)
    if kind == "docx":
        return ExtractedDocument(kind, _docx_text(data), mime=mime)
    if kind == "xlsx":
        rows = _xlsx_rows(data)
        return ExtractedDocument(kind, _rows_text(rows), rows, mime)
    if kind == "csv":
        text = _decode(data)
        try:
            dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(
                text[:2048], delimiters=",;\t"
            )
        except csv.Error:  # a single column has no delimiter to find
            dialect = csv.excel
        rows = [[cell.strip() for cell in row] for row in csv.reader(io.StringIO(text), dialect)]
        return ExtractedDocument(kind, _rows_text(rows), rows, mime)
    if kind == "image":
        return ExtractedDocument(kind, "", mime=mime)
    return ExtractedDocument(kind, _decode(data).strip(), mime=mime)
